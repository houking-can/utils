import json
import os
import re

from docx import Document
from tqdm import tqdm


def iter_files(path):
    """Walk through all files located under a root path."""
    if os.path.isfile(path):
        yield path
    elif os.path.isdir(path):
        for dir_path, _, file_names in os.walk(path):
            for f in file_names:
                yield os.path.join(dir_path, f)
    else:
        raise RuntimeError('Path %s is invalid' % path)


# def format_row(columns):
#     name, appendix, num1, num2 = columns
#
#     if appendix == '-' or appendix == '—':
#         appendix = ''
#
#     if (name.endswith('：') or name.endswith(':')):
#         if num1 in ['-', '', '—', '--'] and num2 in ['-', '', '—', '--']:
#             num1 = None
#             num2 = None
#     else:
#         if num1 in ['-', '', '—', '--']:
#             num1 = '0.00'
#         if num2 in ['-', '', '—', '--']:
#             num2 = '0.00'
#     return {"名称": name,
#             "附注": appendix,
#             "年初至报告期末": num1,
#             "上年年初至报告期末": num2
#             }


def find_unit(text):
    text = re.sub('\s', '', text)
    units = ["单位：元", "单位:元", "单位：美元", "单位:美元", "单位：港元", "单位:港元",
             "单位：欧元", "单位:欧元", "单位：日元", "单位:日元", "单位：英镑", "单位:英镑",
             "单位：加元", "单位:加元", "单位：澳元", "单位:澳元", "单位：卢布", "单位:卢布",
             "单位：韩元", "单位:韩元"]
    for unit in units:
        if unit in text:
            return unit[3:]
    return ''


def find_name(node):
    text_num = 3
    unit = ''
    names = ['母公司资产负债表', '合并资产负债表', '母公司利润表', '合并利润表', '母公司现金流量表', '合并现金流量表']
    for element in node.previous_elements:
        if text_num < 0:
            return None
        if isinstance(element, str):
            element = element.replace('\n', '')
            if element == '':
                continue
            if unit == '':
                unit = find_unit(element)
            for name in names:
                if name in element:
                    if "母公司" in name:
                        return unit, name[3:] + "（母公司）"
                    else:
                        return unit, name[2:] + "（合并）"
            if "利润表" in element:
                return unit, "利润表（合并）"
            if "流量表" in element:
                return unit, "流量表（合并）"
            text_num -= 1


def get_flags(head, name):
    head = head.text.split('\n')
    head = [re.sub('\s', '', each) for each in head]
    if "负债表" in name:
        try:
            ids = [head.index('项目'), head.index('附注'), head.index('期末余额'), head.index('期初余额')]
        except:
            ids = [id for id, val in enumerate(head) if val != '']
    else:
        try:
            ids = [head.index('项目'), head.index('附注'), head.index('本期金额'), head.index('上期金额')]
        except:
            ids = [id for id, val in enumerate(head) if val != '']

    if "负债表" in name:
        tails = ['负债和所有者权益']
    elif "利润表" in name:
        tails = ['稀释每股收益']
    else:
        tails = ['等价物余额', '期末现金']

    return ids, tails


def extract(table, start):
    tab = []
    for i in range(start, len(table.rows)):
        row = []
        for j in range(len(table.columns)):
            row.append(table.cell(i, j).text.replace('\n', '').strip().replace('\t','  '))
        try:
            name, appendix, num1, num2 = row
        except:
            print(file)
            return []
        if appendix == '-' or appendix == '—':
            appendix = ''
        if (name.endswith('：') or name.endswith(':')):
            if num1 in ['-', '', '—', '--'] and num2 in ['-', '', '—', '--']:
                num1 = None
                num2 = None
        else:
            if num1 in ['-', '', '—', '--']:
                num1 = '0.00'
            if num2 in ['-', '', '—', '--']:
                num2 = '0.00'
        row = {"名称": name,
               "附注": appendix,
               "年初至报告期末": num1,
               "上年年初至报告期末": num2
               }
        tab.append(row)
    return tab


def check_end(table, pre_tab):
    head = get_raw(table)
    if "项目附注" in head:
        return True
    if len(table.columns) != 4:
        return True

    flags = ['负债和所有者权益', '稀释每股收益', ['等价物余额', '期末现金']]
    if len(pre_tab)>0:
        tail = pre_tab[-1]['名称']
        if flags[0] in tail or flags[1] in tail:
            return True
        if flags[2][0] in tail and flags[2][1] in tail:
            return True
        if len(pre_tab) > 100:
            return True

    return False


def get_raw(table, start=0):
    row = ''
    for j in range(len(table.columns)):
        row += table.cell(start, j).text
    return re.sub('\s', '', row)


if __name__ == "__main__":
    path = r'C:\Users\Houking\Desktop\error\error\word'
    save_path = r'C:\Users\Houking\Desktop\a'
    files = list(iter_files(path))
    # log = open('doc_log.txt', 'w', encoding='utf-8')
    for file in files:
        # log.write(file+'\n')

        res = dict()
        name = os.path.basename(file)
        name, _ = os.path.splitext(name)
        index = [i.start() for i in re.finditer('-', name)]
        code = name[:index[0]]
        company = name[index[0] + 1:index[1]]
        res[name] = dict()
        res[name]['证券代码'] = code
        res[name]['证券简称'] = company
        res[name]['现金流量表（母公司）'] = {'单位': '', '项目': []}
        res[name]['现金流量表（合并）'] = {'单位': '', '项目': []}
        res[name]['利润表（母公司）'] = {'单位': '', '项目': []}
        res[name]['利润表（合并）'] = {'单位': '', '项目': []}
        res[name]['资产负债表（母公司）'] = {'单位': '', '项目': []}
        res[name]['资产负债表（合并）'] = {'单位': '', '项目': []}

        document = Document(file)
        tables = document.tables
        k = 0
        tmp = []
        while k < len(tables) and len(tmp)<6:
            head = get_raw(tables[k])
            if "项目附注" in head:
                tab = extract(tables[k], start=1)
                k += 1
                while k < len(tables):
                    if check_end(tables[k], tab):
                        break
                    ex_tab = extract(tables[k], start=0)
                    tab.extend(ex_tab)
                    k += 1
                tmp.append(tab)
                continue
            k += 1
        json.dump({"tables": tmp}, open('%s/%s.json' % (save_path, name), 'w', encoding='utf-8'), indent=4,
                  ensure_ascii=False)
        print(tmp)

